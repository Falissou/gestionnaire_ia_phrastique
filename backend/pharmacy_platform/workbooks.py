"""Excel-backed inventory and product-movement repositories."""

from datetime import date, datetime, timezone
from pathlib import Path
from threading import RLock
from uuid import uuid4

from openpyxl import Workbook, load_workbook


def command_signature(command: dict) -> tuple[str, int, str, str | None]:
    expected_date = command.get("expected_date")
    if isinstance(expected_date, datetime):
        expected_date = expected_date.date()
    if isinstance(expected_date, date):
        expected_date = expected_date.isoformat()
    elif expected_date:
        expected_date = str(expected_date)[:10]
    else:
        expected_date = None
    return (
        str(command.get("product_name") or "").strip().casefold(),
        int(command.get("quantity") or 0),
        str(command.get("supplier") or "").strip().casefold(),
        expected_date,
    )


class WorkbookStore:
    STOCK_SHEETS = {
        "Stock": [
            "id",
            "name",
            "category",
            "quantity",
            "min_quantity",
            "unit_price",
            "expiry_date",
            "updated_at",
        ],
        "Commandes": [
            "id",
            "product_name",
            "quantity",
            "supplier",
            "status",
            "order_date",
            "expected_date",
        ],
    }
    MOVEMENT_SHEETS = {
        "Entrees": ["id", "product_id", "product_name", "quantity", "date", "reason", "reference"],
        "Sorties": ["id", "product_id", "product_name", "quantity", "date", "reason", "reference"],
    }

    def __init__(self, directory: Path):
        self.directory = directory
        self.stock_path = directory / "commandes_stock.xlsx"
        self.movement_path = directory / "entrees_sorties.xlsx"
        self.lock = RLock()
        directory.mkdir(parents=True, exist_ok=True)
        self._ensure_workbook(self.stock_path, self.STOCK_SHEETS)
        self._ensure_workbook(self.movement_path, self.MOVEMENT_SHEETS)

    @staticmethod
    def _ensure_workbook(path: Path, sheets: dict[str, list[str]]) -> None:
        if path.exists():
            return
        workbook = Workbook()
        first = True
        for title, headers in sheets.items():
            sheet = workbook.active if first else workbook.create_sheet()
            sheet.title = title
            sheet.append(headers)
            first = False
        workbook.save(path)

    @staticmethod
    def _read(path: Path, sheet_name: str) -> list[dict]:
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)
            return [
                {header: value for header, value in zip(headers, row)}
                for row in rows
                if any(value is not None for value in row)
            ]
        finally:
            workbook.close()

    @staticmethod
    def _write_rows(path: Path, sheet_name: str, rows: list[dict], sheets: dict) -> None:
        workbook = load_workbook(path)
        try:
            sheet = workbook[sheet_name]
            sheet.delete_rows(1, sheet.max_row)
            headers = sheets[sheet_name]
            sheet.append(headers)
            for row in rows:
                sheet.append([row.get(header) for header in headers])
            workbook.save(path)
        finally:
            workbook.close()

    def products(self) -> list[dict]:
        with self.lock:
            return self._read(self.stock_path, "Stock")

    def commands(self) -> list[dict]:
        with self.lock:
            return self._read(self.stock_path, "Commandes")

    def movements(self, kind: str | None = None) -> list[dict]:
        with self.lock:
            sheets = [kind] if kind in self.MOVEMENT_SHEETS else list(self.MOVEMENT_SHEETS)
            return [
                {**row, "type": sheet.lower()}
                for sheet in sheets
                for row in self._read(self.movement_path, sheet)
            ]

    def add_product(self, data: dict) -> dict:
        with self.lock:
            rows = self._read(self.stock_path, "Stock")
            row = {
                "id": str(uuid4()),
                **data,
                "updated_at": datetime.now(timezone.utc).isoformat(),
            }
            rows.append(row)
            self._write_rows(self.stock_path, "Stock", rows, self.STOCK_SHEETS)
            return row

    def add_command(self, data: dict) -> dict:
        with self.lock:
            rows = self._read(self.stock_path, "Commandes")
            signature = command_signature(data)
            if any(
                command_signature(row) == signature and row.get("status") != "Annulée"
                for row in rows
            ):
                raise ValueError("Une commande identique existe déjà.")
            row = {
                "id": str(uuid4()),
                "status": "En attente",
                "order_date": date.today().isoformat(),
                **data,
            }
            rows.append(row)
            self._write_rows(self.stock_path, "Commandes", rows, self.STOCK_SHEETS)
            return row

    def import_commands(self, commands: list[tuple[int, dict]]) -> tuple[list[dict], list[int]]:
        with self.lock:
            rows = self._read(self.stock_path, "Commandes")
            signatures = {
                command_signature(row)
                for row in rows
                if row.get("status") != "Annulée"
            }
            created = []
            duplicates = []
            for source_row, data in commands:
                signature = command_signature(data)
                if signature in signatures:
                    duplicates.append(source_row)
                    continue
                signatures.add(signature)
                command = {
                    "id": str(uuid4()),
                    "status": "En attente",
                    "order_date": date.today().isoformat(),
                    **data,
                }
                rows.append(command)
                created.append(command)
            if created:
                self._write_rows(self.stock_path, "Commandes", rows, self.STOCK_SHEETS)
            return created, duplicates

    def add_movement(self, data: dict) -> dict:
        kind = data["type"]
        sheet_name = "Entrees" if kind == "entree" else "Sorties"
        with self.lock:
            products = self._read(self.stock_path, "Stock")
            product = next((item for item in products if item["id"] == data["product_id"]), None)
            if product is None:
                raise LookupError("Produit introuvable.")
            quantity = data["quantity"]
            if kind == "sortie" and quantity > product["quantity"]:
                raise ValueError("La sortie dépasse la quantité actuellement en stock.")
            product["quantity"] += quantity if kind == "entree" else -quantity
            product["updated_at"] = datetime.now(timezone.utc).isoformat()

            movement = {
                "id": str(uuid4()),
                "product_id": product["id"],
                "product_name": product["name"],
                "quantity": quantity,
                "date": date.today().isoformat(),
                "reason": data["reason"],
                "reference": data.get("reference"),
                "type": kind,
            }
            movements = self._read(self.movement_path, sheet_name)
            movements.append({key: value for key, value in movement.items() if key != "type"})
            self._write_rows(self.stock_path, "Stock", products, self.STOCK_SHEETS)
            self._write_rows(self.movement_path, sheet_name, movements, self.MOVEMENT_SHEETS)
            return movement

    def update_command_status(self, command_id: str, status: str) -> dict | None:
        with self.lock:
            rows = self._read(self.stock_path, "Commandes")
            command = next((row for row in rows if row["id"] == command_id), None)
            if command is None:
                return None
            command["status"] = status
            self._write_rows(self.stock_path, "Commandes", rows, self.STOCK_SHEETS)
            return command
