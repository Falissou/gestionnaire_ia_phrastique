"""Validation helpers for migration spreadsheets and CSV files."""

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import math
import re
import unicodedata
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000

ALIASES = {
    "product_code": {"code", "codeproduit", "productcode", "codemedicament"},
    "name": {
        "nom", "name", "nomproduit", "nommedicament", "nomduproduit",
        "product", "productname", "produit", "article",
    },
    "category": {"categorie", "category"},
    "quantity": {
        "quantite", "qte", "qty", "quantity", "quantitestock",
        "quantitecommandee", "quantiterecue", "quantitesortie", "quantiteconcernee",
        "quantiteaajouter",
    },
    "min_quantity": {"seuil", "seuilalerte", "seuilminimum", "minquantity"},
    "unit_price": {
        "prixunitaire", "prixachatunitaire", "unitprice", "prixunitairefcfa",
        "prixachatunitairefcfa",
    },
    "total_amount": {"montanttotal", "montanttotalfcfa", "totalamount"},
    "expiry_date": {"dateexpiration", "dateperemption", "expirydate"},
    "supplier": {"fournisseur", "supplier", "vendor"},
    "status": {"statut", "statutcommande", "status"},
    "source_status": {"statutstock"},
    "source_order_number": {"ncommande", "numerocommande", "ordernumber"},
    "responsible": {"responsable", "responsible"},
    "comments": {"commentaire", "commentaires", "comment", "comments", "notes"},
    "entry_date": {"dateentree", "entrydate"},
    "order_date": {"datecommande", "orderdate"},
    "expected_date": {
        "dateprevue", "dateprevuedelivraison", "datedelivraisonprevue",
        "datelivraisonprevue", "expecteddate", "deliverydate",
    },
    "date": {"date", "dateentree", "dateentreeprevue", "datesortie"},
    "reason": {"motif", "raison", "reason"},
    "reference": {
        "reference", "numerobonlivraison", "nbonlivraison",
        "serviceclient", "service", "observation", "commentaires",
    },
}


def _normalise_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]", "", text)


def _parse_date(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, date_format).date().isoformat()
        except ValueError:
            continue
    raise ValueError("date invalide (formats acceptés : AAAA-MM-JJ ou JJ/MM/AAAA)")


def _number(value: object, *, integer: bool = False, allow_empty: bool = False):
    if value is None or value == "":
        if allow_empty:
            return None
        raise ValueError("valeur numérique manquante")
    if isinstance(value, bool):
        raise ValueError("valeur numérique invalide")
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError("valeur numérique invalide") from exc
    if not math.isfinite(number) or number < 0:
        raise ValueError("la valeur ne peut pas être négative")
    if integer:
        if not number.is_integer():
            raise ValueError("la quantité doit être un nombre entier")
        return int(number)
    return number


def read_tabular_file(filename: str, content: bytes) -> list[tuple[int, list[object]]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="La taille maximale du fichier est de 5 Mo.")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if extension == "csv":
            text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(StringIO(text), dialect))
        elif extension == "xlsx":
            workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            raise HTTPException(status_code=415, detail="Utilisez un fichier .xlsx ou .csv.")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="Le fichier CSV doit être encodé en UTF-8.") from exc
    except (csv.Error, BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=400, detail="Le fichier Excel ou CSV est illisible ou invalide.") from exc
    if not rows:
        raise HTTPException(status_code=400, detail="Le fichier ne contient aucune ligne.")
    if len(rows) > MAX_IMPORT_ROWS + 1:
        raise HTTPException(
            status_code=413,
            detail=f"Le fichier dépasse la limite de {MAX_IMPORT_ROWS} lignes.",
        )
    return [(number, list(row)) for number, row in enumerate(rows, start=1)]


def parse_import_rows(target: str, rows: list[tuple[int, list[object]]]) -> list[dict]:
    if not rows:
        raise HTTPException(status_code=400, detail="Le fichier ne contient aucune donnée à importer.")
    headers = rows[0][1]
    columns: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = _normalise_header(header)
        for field, aliases in ALIASES.items():
            if normalized in aliases and field not in columns:
                columns[field] = index
    required = {
        "product": {"name", "quantity"},
        "command": {"name", "quantity", "supplier"},
        "entree": {"quantity", "date"},
        "sortie": {"quantity", "date"},
    }[target]
    if target in ("entree", "sortie") and not ({"name", "product_code"} & columns.keys()):
        raise HTTPException(status_code=400, detail="Une colonne Code Produit ou Nom Médicament est obligatoire.")
    missing = required - columns.keys()
    if missing:
        labels = {
            "name": "nom/produit",
            "quantity": "quantité",
            "supplier": "fournisseur",
            "date": "date",
        }
        raise HTTPException(
            status_code=400,
            detail="Colonnes obligatoires manquantes : " + ", ".join(labels.get(field, field) for field in sorted(missing)),
        )

    records = []
    for row_number, values in rows[1:]:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        raw = {
            field: values[index] if index < len(values) else None
            for field, index in columns.items()
        }
        try:
            quantity = _number(raw.get("quantity"), integer=True)
            if target != "product" and quantity < 1:
                raise ValueError("la quantité doit être supérieure à zéro")
            if target == "product":
                name = str(raw.get("name") or "").strip()
                if not name:
                    raise ValueError("nom du produit manquant")
                record = {
                    "product_code": str(raw.get("product_code") or "").strip() or None,
                    "name": name,
                    "category": str(raw.get("category") or "").strip() or "Médicament",
                    "supplier": str(raw.get("supplier") or "").strip() or None,
                    "quantity": quantity,
                    "min_quantity": _number(raw.get("min_quantity"), integer=True, allow_empty=True),
                    "unit_price": _number(raw.get("unit_price"), allow_empty=True),
                    "expiry_date": _parse_date(raw.get("expiry_date")),
                    "entry_date": _parse_date(raw.get("entry_date")),
                    "source_status": str(raw.get("source_status") or "").strip() or None,
                    "comments": str(raw.get("comments") or "").strip() or None,
                }
                record["min_quantity"] = record["min_quantity"] if record["min_quantity"] is not None else 5
                record["unit_price"] = record["unit_price"] if record["unit_price"] is not None else 0
            elif target == "command":
                status = str(raw.get("status") or "").strip().casefold()
                status_map = {
                    "annulee": "Annulée",
                    "annulée": "Annulée",
                    "commandee": "Commandée",
                    "validée": "Commandée",
                    "validee": "Commandée",
                    "en cours de livraison": "Commandée",
                    "reçue": "Reçue",
                    "recue": "Reçue",
                    "livree": "Reçue",
                    "livrée": "Reçue",
                }
                record = {
                    "source_order_number": str(raw.get("source_order_number") or "").strip() or None,
                    "product_name": str(raw.get("name") or "").strip(),
                    "product_code": str(raw.get("product_code") or "").strip() or None,
                    "quantity": quantity,
                    "supplier": str(raw.get("supplier") or "").strip(),
                    "unit_price": _number(raw.get("unit_price"), allow_empty=True),
                    "total_amount": _number(raw.get("total_amount"), allow_empty=True),
                    "responsible": str(raw.get("responsible") or "").strip() or None,
                    "comments": str(raw.get("comments") or "").strip() or None,
                    "command_status": status_map.get(status, "En attente"),
                    "order_date": _parse_date(raw.get("order_date")) or date.today().isoformat(),
                    "expected_date": _parse_date(raw.get("expected_date")),
                }
                if not record["product_name"] or not record["supplier"]:
                    raise ValueError("produit ou fournisseur manquant")
            else:
                name = str(raw.get("name") or "").strip()
                product_code = str(raw.get("product_code") or "").strip() or None
                if not name and not product_code:
                    raise ValueError("code produit ou nom manquant")
                movement_date = _parse_date(raw.get("date"))
                if movement_date is None:
                    raise ValueError("date manquante")
                record = {
                    "product_code": product_code,
                    "product_name": name,
                    "quantity": quantity,
                    "date": movement_date,
                    "reason": str(raw.get("reason") or "").strip() or (
                        "Entrée importée" if target == "entree" else "Sortie importée"
                    ),
                    "reference": str(raw.get("reference") or "").strip() or None,
                }
            records.append({
                "source_row": row_number,
                **record,
                "status": "ready",
                "error": None,
            })
        except (TypeError, ValueError) as exc:
            records.append({
                "source_row": row_number,
                **{field: raw.get(field) for field in columns},
                "status": "invalid",
                "error": str(exc),
            })
    if not records:
        raise HTTPException(status_code=400, detail="Le fichier ne contient aucune ligne à traiter.")
    return records
