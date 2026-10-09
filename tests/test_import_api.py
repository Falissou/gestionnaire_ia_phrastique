from io import BytesIO
import tempfile
import unittest
import unicodedata
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from pharmacy_platform import main
from pharmacy_platform.sqlite_store import SQLiteStore


class ImportApiTests(unittest.TestCase):
    def test_updated_reference_workbooks_are_listed_even_with_decomposed_accents(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            filenames = [
                "Stock_Medicaments.xlsx",
                "Commandes_type.xlsx",
                "Entrée_type.xlsx",
                "Sortis_type.xlsx",
            ]
            for filename in filenames:
                (data_dir / unicodedata.normalize("NFD", filename)).touch()

            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with (
                    patch.object(main.settings, "data_dir", data_dir),
                    TestClient(main.app) as client,
                ):
                    references = client.get("/api/import/references")
                self.assertEqual(
                    {unicodedata.normalize("NFC", item["filename"]): item["target"] for item in references.json()},
                    {
                        "Stock_Medicaments.xlsx": "product",
                        "Commandes_type.xlsx": "command",
                        "Entrée_type.xlsx": "entree",
                        "Sortis_type.xlsx": "sortie",
                    },
                )
                self.assertTrue(all(item["default"] for item in references.json()))
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_import_automatically_uses_reference_for_selected_data_type(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_dir = Path(directory)
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "Code Produit", "Nom Médicament", "Catégorie", "Fournisseur",
                "Quantité Stock", "Seuil Alerte", "Prix Unitaire (FCFA)",
                "Date Entrée", "Date Expiration", "Statut Stock", "Commentaires",
            ])
            sheet.append([
                "MED001", "Paracétamol", "Antalgique", "Pharma Plus", 19, 5,
                500, "03/09/2026", "21/01/2027", "Disponible", "Stock initial",
            ])
            workbook.save(data_dir / "Stock_Medicaments.xlsx")

            store = SQLiteStore(data_dir / "test.sqlite3")
            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with (
                    patch.object(main, "store", store),
                    patch.object(main.settings, "data_dir", data_dir),
                    TestClient(main.app) as client,
                ):
                    preview = client.post(
                        "/api/import/preview",
                        data={"target": "product"},
                    )
                self.assertEqual(preview.status_code, 200, preview.text)
                self.assertEqual(preview.json()["filename"], "Stock_Medicaments.xlsx")
                self.assertEqual(preview.json()["ready_count"], 1)
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_import_requires_a_valid_one_time_preview(self) -> None:
        csv_content = (
            "Code Produit,Nom Médicament,Catégorie,Fournisseur,Quantité Stock,"
            "Seuil Alerte,Prix Unitaire (FCFA),Date Entrée,Date Expiration,"
            "Statut Stock,Commentaires\n"
            "MED001,Paracétamol,Antalgique,Pharma Plus,19,5,500,"
            "03/09/2026,21/01/2027,Disponible,Stock initial\n"
        )
        with tempfile.TemporaryDirectory() as directory:
            store = SQLiteStore(Path(directory) / "test.sqlite3")
            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with (
                    patch.object(main, "store", store),
                    patch.object(main.settings, "data_dir", Path(directory)),
                    TestClient(main.app) as client,
                ):
                    preview = client.post(
                        "/api/import/preview",
                        data={"target": "product"},
                        files={"file": ("stock.csv", csv_content, "text/csv")},
                    )
                    self.assertEqual(preview.status_code, 200, preview.text)
                    result = preview.json()
                    self.assertEqual(result["ready_count"], 1)
                    self.assertTrue(result["preview_id"])
                    self.assertEqual(store.products(), [])

                    imported = client.post(
                        "/api/import/confirm",
                        data={"preview_id": result["preview_id"]},
                    )
                    self.assertEqual(imported.status_code, 200, imported.text)
                    self.assertEqual(imported.json()["created_count"], 1)
                    self.assertEqual(store.products()[0]["product_code"], "MED001")
                    product = store.products()[0]
                    self.assertEqual(product["supplier"], "Pharma Plus")
                    self.assertEqual(product["entry_date"], "2026-09-03")
                    self.assertEqual(product["source_status"], "Disponible")
                    self.assertEqual(product["comments"], "Stock initial")

                    replay = client.post(
                        "/api/import/confirm",
                        data={"preview_id": result["preview_id"]},
                    )
                    self.assertEqual(replay.status_code, 409)
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_updated_command_import_keeps_validated_status(self) -> None:
        csv_content = (
            "N° Commande,Date Commande,Fournisseur,Code Produit,Nom Médicament,"
            "Quantité Commandée,Prix Unitaire (FCFA),Montant Total (FCFA),"
            "Date Livraison Prévue,Statut Commande,Responsable,Commentaires\n"
            "CMD0001,06/01/2026,Fournisseur 2,MED001,Paracétamol,52,550,28600,"
            "20/01/2026,Validée,Responsable 2,À vérifier\n"
        )
        with tempfile.TemporaryDirectory() as directory:
            store = SQLiteStore(Path(directory) / "test.sqlite3")
            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with patch.object(main, "store", store), TestClient(main.app) as client:
                    preview = client.post(
                        "/api/import/preview",
                        data={"target": "command"},
                        files={"file": ("Commandes_type.csv", csv_content, "text/csv")},
                    )
                    self.assertEqual(preview.status_code, 200, preview.text)
                    result = preview.json()
                    self.assertEqual(result["ready_count"], 1)
                    self.assertEqual(result["rows"][0]["command_status"], "Commandée")
                    imported = client.post(
                        "/api/import/confirm",
                        data={"preview_id": result["preview_id"]},
                    )
                    self.assertEqual(imported.status_code, 200, imported.text)
                    command = store.commands()[0]
                    self.assertEqual(command["status"], "Commandée")
                    self.assertEqual(command["order_date"], "2026-01-06")
                    self.assertEqual(command["expected_date"], "2026-01-20")
                    self.assertEqual(command["source_order_number"], "CMD0001")
                    self.assertEqual(command["product_code"], "MED001")
                    self.assertEqual(command["unit_price"], 550)
                    self.assertEqual(command["total_amount"], 28600)
                    self.assertEqual(command["responsible"], "Responsable 2")
                    self.assertEqual(command["comments"], "À vérifier")
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_command_page_import_accepts_updated_headers_and_preserves_status(self) -> None:
        csv_content = (
            "N° Commande,Date Commande,Fournisseur,Code Produit,Nom Médicament,"
            "Quantité Commandée,Date Livraison Prévue,Statut Commande\n"
            "CMD0002,11/01/2026,Fournisseur 3,MED002,Amoxicilline,54,"
            "25/01/2026,En cours de livraison\n"
        )
        with tempfile.TemporaryDirectory() as directory:
            store = SQLiteStore(Path(directory) / "test.sqlite3")
            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with patch.object(main, "store", store), TestClient(main.app) as client:
                    preview = client.post(
                        "/api/commands/import/preview",
                        files={"file": ("Commandes_type.csv", csv_content, "text/csv")},
                    )
                    self.assertEqual(preview.status_code, 200, preview.text)
                    result = preview.json()
                    self.assertEqual(result["ready_count"], 1)
                    self.assertEqual(result["rows"][0]["command_status"], "Commandée")
                    imported = client.post(
                        "/api/commands/import",
                        data={"preview_id": result["preview_id"]},
                    )
                    self.assertEqual(imported.status_code, 200, imported.text)
                    command = store.commands()[0]
                    self.assertEqual(command["status"], "Commandée")
                    self.assertEqual(command["order_date"], "2026-01-11")
                    self.assertEqual(command["expected_date"], "2026-01-25")
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_movement_registers_create_products_and_adjust_stock(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SQLiteStore(Path(directory) / "test.sqlite3")
            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with patch.object(main, "store", store), TestClient(main.app) as client:
                    entry = client.post("/api/movements", json={
                        "type": "entree",
                        "quantity": 8,
                        "reason": "Réception",
                        "product_code": "MED004",
                        "product_name": "Vitamine C",
                    })
                    self.assertEqual(entry.status_code, 201, entry.text)
                    product_id = entry.json()["product_id"]
                    self.assertEqual(store.products()[0]["quantity"], 8)

                    exit_response = client.post("/api/movements", json={
                        "product_id": product_id,
                        "type": "sortie",
                        "quantity": 3,
                        "reason": "Vente",
                    })
                    self.assertEqual(exit_response.status_code, 201, exit_response.text)
                    self.assertEqual(store.products()[0]["quantity"], 5)
                    self.assertEqual(len(client.get("/api/movements?kind=entree").json()), 1)
                    self.assertEqual(len(client.get("/api/movements?kind=sortie").json()), 1)
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_purchase_order_workbook_imports_as_received_stock_entry(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SQLiteStore(Path(directory) / "test.sqlite3")
            product = store.add_product({
                "product_code": "MED015",
                "name": "Produit en rupture",
                "category": "Médicament",
                "quantity": 0,
                "min_quantity": 5,
                "unit_price": 250,
                "expiry_date": None,
            })
            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with patch.object(main, "store", store), TestClient(main.app) as client:
                    order = client.post("/api/commands/place-order", json={
                        "supplier": "Fournisseur test",
                        "expected_date": "2026-10-10",
                        "responsible": "Équipe achats",
                        "comments": "Réapprovisionnement",
                        "lines": [{
                            "product_id": product["id"],
                            "quantity": 5,
                            "unit_price": 250,
                        }],
                    })
                    self.assertEqual(order.status_code, 200, order.text)
                    workbook = load_workbook(BytesIO(order.content))
                    try:
                        self.assertEqual(
                            workbook.sheetnames,
                            ["Import entrée", "Bon de commande"],
                        )
                        entry_sheet = workbook["Import entrée"]
                        self.assertEqual(
                            [entry_sheet.cell(1, column).value for column in range(1, 7)],
                            [
                                "Code Produit",
                                "Nom Médicament",
                                "Quantité à Ajouter",
                                "Motif",
                                "Date Entrée",
                                "Référence",
                            ],
                        )
                        self.assertEqual(entry_sheet["A2"].value, "MED015")
                        self.assertEqual(entry_sheet["C2"].value, 5)
                        self.assertIsNone(entry_sheet["E2"].value)
                        self.assertEqual(entry_sheet["F2"].value, store.commands()[0]["source_order_number"])
                        entry_sheet["E2"] = "06/10/2026"
                        output = BytesIO()
                        workbook.save(output)
                    finally:
                        workbook.close()

                    preview = client.post(
                        "/api/import/preview",
                        data={"target": "entree", "adjust_stock": "true"},
                        files={
                            "file": (
                                "bon_commande.xlsx",
                                output.getvalue(),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            ),
                        },
                    )
                    self.assertEqual(preview.status_code, 200, preview.text)
                    result = preview.json()
                    self.assertEqual(result["ready_count"], 1)
                    self.assertEqual(result["invalid_count"], 0)
                    self.assertEqual(result["rows"][0]["product_code"], "MED015")
                    self.assertEqual(result["rows"][0]["date"], "2026-10-06")
                    imported = client.post(
                        "/api/import/confirm",
                        data={"preview_id": result["preview_id"]},
                    )
                    self.assertEqual(imported.status_code, 200, imported.text)
                    self.assertEqual(store.products()[0]["quantity"], 5)
                    movement = store.movements("entree")[0]
                    self.assertEqual(movement["quantity"], 5)
                    self.assertEqual(movement["reference"], store.commands()[0]["source_order_number"])
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)

    def test_validated_command_export_contains_only_commanded_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            store = SQLiteStore(Path(directory) / "test.sqlite3")
            product = store.add_product({
                "product_code": "MED001",
                "name": "=Produit test",
                "category": "Médicament",
                "quantity": 4,
                "min_quantity": 5,
                "unit_price": 500,
                "expiry_date": None,
            })
            validated = store.add_command({
                "product_name": product["name"],
                "quantity": 6,
                "supplier": "Fournisseur",
                "expected_date": "2026-10-10",
                "source_order_number": "CMD-REF",
                "unit_price": 125.5,
                "total_amount": 753,
                "responsible": "Équipe achats",
                "comments": "Prioritaire",
            })
            store.update_command_status(validated["id"], "Commandée")
            store.add_command({
                "product_name": product["name"],
                "quantity": 7,
                "supplier": "Autre fournisseur",
                "expected_date": None,
            })

            main.app.dependency_overrides[main.require_auth] = lambda: None
            try:
                with patch.object(main, "store", store), TestClient(main.app) as client:
                    response = client.get("/api/commands/export/validated")
                self.assertEqual(response.status_code, 200, response.text)
                self.assertIn("attachment; filename=\"commandes_validees.xlsx\"", response.headers["content-disposition"])
                workbook = load_workbook(BytesIO(response.content), data_only=True)
                try:
                    sheet = workbook.active
                    self.assertEqual(sheet.max_row, 2)
                    self.assertEqual(sheet.cell(2, 4).value, "MED001")
                    self.assertEqual(sheet.cell(2, 5).value, "=Produit test")
                    self.assertEqual(sheet.cell(2, 5).data_type, "s")
                    self.assertEqual(sheet.cell(2, 1).value, "CMD-REF")
                    self.assertEqual(sheet.cell(2, 7).value, 125.5)
                    self.assertEqual(sheet.cell(2, 8).value, 753)
                    self.assertEqual(sheet.cell(2, 10).value, "Commandée")
                    self.assertEqual(sheet.cell(2, 12).value, "Prioritaire")
                finally:
                    workbook.close()
            finally:
                main.app.dependency_overrides.pop(main.require_auth, None)


if __name__ == "__main__":
    unittest.main()
